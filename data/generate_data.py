"""
Generates a FICTIONAL dataset for the delivery-date-change dashboard.
- 3,200 orders (Auftraege), ~1,760 delivery date changes (Terminaenderungen)
- Deliberately injects typical manual-entry data quality issues (see docs/data_quality_rules.md)
Product, plant and dealer names are invented. No real company or customer data.
Usage: python generate_data.py   (requires pandas, openpyxl)
"""
import random, datetime as dt, pandas as pd
random.seed(42)
D=dt.date
start, end = D(2025,1,6), D(2026,9,25)
models = [("Van L","Werk A",.35),("Van L Fahrgestell","Werk B",.15),("eVan L","Werk A",.08),
          ("Van M","Werk C",.17),("eVan M","Werk C",.05),("MPV","Werk C",.12),("MPV Asia","Werk D",.08)]
markets = [("Deutschland",.30),("Frankreich",.12),("Italien",.08),("Spanien",.08),("Niederlande",.06),("Großbritannien",.08),
           ("Polen",.05),("Algerien",.04),("Türkei",.04),("China",.10),("VAE",.05)]
reasons = [("Teilemangel",.30),("Produktionsstörung",.15),("Kundenänderung",.18),("Logistik/Transport",.15),
           ("Kapazität/Einplanung",.12),("Zoll/Export",.05),("Qualitätsnacharbeit",.05)]
ktyp=[("Flotte",.45),("Gewerbe",.40),("Privat",.15)]
def pick(lst): 
    r=random.random(); c=0
    for item in lst:
        c+=item[-1]
        if r<=c: return item
    return lst[-1]
orders=[]; changes=[]; cid=1
for i in range(1,3201):
    od = start + dt.timedelta(days=random.randint(0,(end-start).days))
    m,w,_ = pick(models)
    if m=="MPV Asia": mk="China"
    else:
        mk=pick(markets)[0]
        if mk=="China": mk=random.choice(["Deutschland","Frankreich"])
    am=f"AM-{mk[:2].upper()}-{random.randint(1,25):03d}"
    first = od + dt.timedelta(days=random.randint(60,120))
    cur=first
    p_change = 0.38 + (0.25 if mk=="China" else 0) + (0.10 if m.startswith("e") else 0)
    n = 0
    if random.random()<p_change:
        n = 1 if random.random()<.7 else (2 if random.random()<.8 else 3)
    cdate = od
    for k in range(n):
        cdate = cdate + dt.timedelta(days=random.randint(5,40))
        if cdate>=cur: cdate=cur-dt.timedelta(days=3)
        rs = pick(reasons)[0]
        if mk=="China" and random.random()<.4: rs=random.choice(["Zoll/Export","Logistik/Transport"])
        if random.random()<.08: shift=-random.randint(3,14)
        else:
            base={"Teilemangel":21,"Produktionsstörung":10,"Kundenänderung":14,"Logistik/Transport":9,"Kapazität/Einplanung":12,"Zoll/Export":18,"Qualitätsnacharbeit":7}[rs]
            shift=max(1,int(random.gauss(base,base*0.5)))
            if mk=="China": shift=int(shift*1.5)
        new = cur+dt.timedelta(days=shift)
        inf = "Ja" if random.random()<(0.78 if mk!="China" else 0.6) else "Nein"
        erf = "Schnittstelle" if random.random()<.65 else "Manuell"
        changes.append(dict(Aenderungs_ID=f"TA-{cid:05d}",Auftragsnr=f"VAN-{i:06d}",Aenderungsdatum=cdate,
                            Alter_Termin=cur,Neuer_Termin=new,Verschiebung_Tage=(new-cur).days,Grund=rs,
                            Kunde_informiert=inf,Erfasst_ueber=erf))
        cid+=1; cur=new
    today=D(2026,10,1)
    actual=None; status="Offen"
    if cur<today:
        actual = cur + dt.timedelta(days=max(-2,int(random.gauss(2,4))))
        status="Ausgeliefert" if actual<today else "Offen"
        if status=="Offen": actual=None
    orders.append(dict(Auftragsnr=f"VAN-{i:06d}",Auftragsdatum=od,Modell=m,Werk=w,Markt=mk,Absatzmittler=am,
                       Kundentyp=pick(ktyp)[0],Erster_Liefertermin=first,Aktueller_Liefertermin=cur,
                       Tatsaechlicher_Liefertermin=actual,Status=status))
O=pd.DataFrame(orders); C=pd.DataFrame(changes)
# ---- inject data quality issues (manual entries mostly) ----
man = C.index[C.Erfasst_ueber=="Manuell"].tolist()
random.shuffle(man)
log={}
a=man[:45];  C.loc[a,"Grund"]=None; log["Grund leer"]=len(a)
b=man[45:85]; C.loc[b,"Kunde_informiert"]=[random.choice(["yes","vielleicht","","j","no"]) for _ in b]; log["Kunde_informiert ungültig"]=len(b)
c=man[85:110]; C.loc[c,"Verschiebung_Tage"]=C.loc[c,"Verschiebung_Tage"]+[random.choice([-7,7,10,-10,30]) for _ in c]; log["Verschiebung passt nicht zu Daten"]=len(c)
d=man[110:125]
for ix in d: C.at[ix,"Neuer_Termin"]=C.at[ix,"Alter_Termin"]-dt.timedelta(days=400)
log["Neuer Termin unplausibel (>1 Jahr zurück)"]=len(d)
dup=C.loc[man[125:155]].copy(); C=pd.concat([C,dup]).sort_values(["Aenderungsdatum","Aenderungs_ID"]).reset_index(drop=True); log["Doppelte Änderungen"]=len(dup)
oi=random.sample(range(len(O)),60)
variants={"Frankreich":["France","frankreich"],"Deutschland":["DE","Germany"],"Spanien":["Spain"],"Italien":["Italy"],"Großbritannien":["UK","Grossbritannien"],"China":["CN","china"]}
cnt=0
for ix in oi:
    mk=O.at[ix,"Markt"]
    if mk in variants: O.at[ix,"Markt"]=random.choice(variants[mk]); cnt+=1
log["Markt-Schreibweisen uneinheitlich"]=cnt
e=random.sample(range(len(O)),20); O.loc[e,"Absatzmittler"]=None; log["Absatzmittler leer"]=20
print(len(O),"orders",len(C),"changes"); print(log)
import os
out=os.path.join(os.path.dirname(os.path.abspath(__file__)),"Liefertermin_Prototyp_Daten.xlsx")
with pd.ExcelWriter(out, engine="openpyxl") as xw:
    pd.DataFrame({"Hinweis":["FIKTIVE DATEN – selbst generiert für einen Übungs-Prototyp. Keine echten Auftrags- oder Kundendaten.",
                             "Tabellen: Auftraege (1 Zeile pro Auftrag), Terminaenderungen (1 Zeile pro Liefertermin-Änderung).",
                             "Verknüpfung: Auftraege[Auftragsnr] 1 → * Terminaenderungen[Auftragsnr]",
                             "Datenstand: 01.10.2026"]}).to_excel(xw,sheet_name="README",index=False)
    O.to_excel(xw,sheet_name="Auftraege",index=False); C.to_excel(xw,sheet_name="Terminaenderungen",index=False)
from openpyxl import load_workbook
from openpyxl.styles import Font
from openpyxl.worksheet.table import Table, TableStyleInfo
wb=load_workbook(out)
for ws in wb:
    for row in ws.iter_rows():
        for cell in row:
            cell.font=Font(name="Arial",size=10,bold=(cell.row==1))
            if isinstance(cell.value,(dt.date,dt.datetime)): cell.number_format="DD.MM.YYYY"
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width=max(12,min(60,max(len(str(c.value or "")) for c in col[:200])+2))
    if ws.title!="README":
        ref=f"A1:{ws.cell(1,ws.max_column).column_letter}{ws.max_row}"
        t=Table(displayName=ws.title,ref=ref); t.tableStyleInfo=TableStyleInfo(name="TableStyleMedium2",showRowStripes=True)
        ws.add_table(t); ws.freeze_panes="A2"
wb.save(out)
